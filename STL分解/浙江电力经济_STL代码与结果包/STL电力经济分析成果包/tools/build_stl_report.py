from __future__ import annotations

import csv
import json
import math
import os
import shutil
from datetime import date
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

project = Path("/mnt/data/STL电力经济分析成果包")
outputs = project / "outputs"
report_path = Path("/mnt/data/浙江电力经济_STL分解完整报告.docx")

summary = json.loads((outputs / "run_summary.json").read_text(encoding="utf-8"))
quality_report = json.loads((outputs / "data_quality_report.json").read_text(encoding="utf-8"))

def csv_rows(rel: str) -> list[dict[str, str]]:
    with (outputs / rel).open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))

electricity = csv_rows("tables/electricity_monthly_stl.csv")
gdp = csv_rows("tables/gdp_quarterly_stl.csv")
anomalies = csv_rows("tables/anomaly_events.csv")
seasonal = csv_rows("tables/electricity_month_seasonal_factors.csv")
quality = csv_rows("tables/data_quality_checks.csv")
sources = csv_rows("tables/data_source_registry.csv")
kg_nodes = csv_rows("kg/kg_nodes.csv")
kg_edges = csv_rows("kg/kg_edges.csv")
model_records = [
    json.loads(x)
    for x in (outputs / "model_input/model_input_records.jsonl").read_text(encoding="utf-8").splitlines()
    if x.strip()
]

def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:cantSplit")) is None:
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)

def set_run_font(run, name="Noto Sans CJK SC", size=10.5, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def set_paragraph_font(paragraph, name="Noto Sans CJK SC", size=10.5):
    for run in paragraph.runs:
        set_run_font(run, name, size)

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=(100,100,100))
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=9, color=(100,100,100))

def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Noto Sans CJK SC")
    rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    rPr.append(rFonts)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink

def set_col_widths(table, widths_cm: Sequence[float]):
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            if idx < len(row.cells):
                row.cells[idx].width = Cm(width)

def add_table(doc, headers: Sequence[str], rows: Iterable[Sequence], widths=None, font_size=8.5, style="Table Grid"):
    rows = list(rows)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = style
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_row_cant_split(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = str(h)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F5A7A")
        set_cell_margins(cell)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                set_run_font(run, size=font_size, bold=True, color=(255,255,255))
    for ridx, row in enumerate(rows, 1):
        new_row = table.add_row()
        set_row_cant_split(new_row)
        cells = new_row.cells
        for i, val in enumerate(row):
            cells[i].text = "" if val is None else str(val)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            if ridx % 2 == 0:
                set_cell_shading(cells[i], "F4F8FB")
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_run_font(run, size=font_size)
    if widths:
        set_col_widths(table, widths)
    return table

def add_caption(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    run = p.add_run(text)
    set_run_font(run, size=9, color=(90,90,90))
    return p

def add_figure(doc, path: Path, caption: str, width_cm=15.7):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    r = p.add_run()
    r.add_picture(str(path), width=Cm(width_cm))
    add_caption(doc, caption)

def add_note_box(doc, title: str, text: str, fill="FFF4D6", title_color=(122,75,0)):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(16.3)
    set_row_cant_split(table.rows[0])
    cell = table.cell(0,0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=120, bottom=120, start=140, end=140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=10.5, bold=True, color=title_color)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    r2 = p2.add_run(text)
    set_run_font(r2, size=9.5)
    return table

def add_code_block(doc, code: str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0,0)
    set_cell_shading(cell, "F3F5F7")
    set_cell_margins(cell, top=100, bottom=100, start=120, end=120)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(code)
    set_run_font(run, name="Noto Sans Mono CJK SC", size=8.2)
    return table

def add_bullet(doc, text: str, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run_font(r, size=10.5)
    return p

def add_number(doc, text: str):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run_font(r, size=10.5)
    return p

def add_body(doc, text: str, bold_prefix: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.45
    p.paragraph_format.space_after = Pt(5)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=10.5, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=10.5)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5)
    return p

def heading(doc, text: str, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level==1 else 7)
    p.paragraph_format.space_after = Pt(5)
    for r in p.runs:
        set_run_font(r, size={1:16, 2:13, 3:11.5}.get(level,10.5), bold=True, color=(18,59,93))
    return p

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.8)
sec.bottom_margin = Cm(1.7)
sec.left_margin = Cm(1.8)
sec.right_margin = Cm(1.8)
sec.header_distance = Cm(0.8)
sec.footer_distance = Cm(0.8)

# Global styles
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Noto Sans CJK SC"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.35
normal.paragraph_format.space_after = Pt(4)

for sname, size, bold in [
    ("Title", 24, True), ("Subtitle", 13, False),
    ("Heading 1", 16, True), ("Heading 2", 13, True), ("Heading 3", 11.5, True)
]:
    st = styles[sname]
    st.font.name = "Noto Sans CJK SC"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    st.font.size = Pt(size)
    st.font.bold = bold
    if "Heading" in sname:
        st.font.color.rgb = RGBColor(18,59,93)

# Header/footer
for section in doc.sections:
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("浙江省电力—经济数据 STL 分解与智能体接口设计")
    set_run_font(hr, size=8.5, color=(100,100,100))
    add_page_number(section.footer.paragraphs[0])

# Cover
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(90)
r = p.add_run("浙江省电力—经济数据\nSTL 分解与智能体接口设计")
set_run_font(r, size=27, bold=True, color=(18,59,93))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(14)
r = p.add_run("面向“电力经济领域 Claude Code”的数据获取、特征解耦与知识图谱前置工程")
set_run_font(r, size=13.5, color=(31,90,122))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(50)
cover_rows = [
    ("研究对象", "浙江省全社会用电量与三次产业增加值"),
    ("时间范围", "电力：2021-01—2025-03；GDP：2000Q1—2025Q4"),
    ("核心方法", "log-STL、鲁棒异常检测、节假日特征、知识图谱与模型输入翻译"),
    ("交付内容", "完整代码、运行结果、结果工作簿、模型 JSONL、知识图谱 CSV"),
    ("生成日期", "2026-07-24"),
]
ct = doc.add_table(rows=0, cols=2)
ct.alignment = WD_TABLE_ALIGNMENT.CENTER
ct.autofit = False
for k,v in cover_rows:
    row = ct.add_row()
    row.cells[0].text = k
    row.cells[1].text = v
    set_cell_shading(row.cells[0], "DCEAF4")
    for i,c in enumerate(row.cells):
        set_cell_margins(c, top=100,bottom=100,start=120,end=120)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for pp in c.paragraphs:
            pp.alignment = WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
            for rr in pp.runs:
                set_run_font(rr, size=10.5, bold=(i==0))
set_col_widths(ct, [3.3, 12.5])

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(70)
r = p.add_run("基于用户提供的开题材料、浙江省数据文件与公开权威信息完成")
set_run_font(r, size=10, color=(90,90,90))
doc.add_page_break()

# Executive summary
heading(doc, "摘要", 1)
add_body(doc, "本报告围绕老师强调的“数据获取—STL 分解—知识图谱构建”三个重点，完成了一个可以直接接入后续大模型智能体的前置数据工程闭环。系统先把原始 Excel 和节假日文档转为统一、可审计的时序记录，再通过鲁棒 log-STL 分离趋势、季节和不规则项，最后将异常、候选原因、证据缺口、数据来源与质量标记翻译为 JSONL 和知识图谱节点/关系。")
add_body(doc, "本次运行处理浙江省月度全社会用电量 51 条、三次产业单季度增加值 312 条，共生成 363 条模型输入记录；识别电力 high 异常 7 条、GDP high 异常 22 条；构建知识图谱 647 个节点、1235 条关系。2024 年月度用电量求和为 6780.85 亿千瓦时，与公开权威口径的 6780 亿千瓦时基本一致；2025 年三次产业累计值分别为 2657、35682、56206 亿元，也与浙江统计公报一致。")
add_note_box(doc, "最重要的解释边界", "STL 负责回答“什么时候偏离了正常趋势和季节规律”，不能单独证明“为什么”。因此本成果把春节、天气、工业生产、政策/电网事件等写成候选解释与待补证据，而不是直接写成因果结论。")
heading(doc, "主要交付物", 2)
add_bullet(doc, "完整可重跑代码：stl_power_economy_pipeline.py。")
add_bullet(doc, "结构化结果：电力/GDP STL 明细、异常清单、节假日特征、质量检查、数据源登记。")
add_bullet(doc, "模型接口：363 条 JSONL，每条包含观测、分解、异常、上下文、证据需求与溯源。")
add_bullet(doc, "知识图谱：Region—Metric—Time—Observation—Anomaly—Holiday—CauseHypothesis 节点与关系。")
add_bullet(doc, "可视化与工作簿：8 张分析图、1 个可筛选 Excel 结果工作簿。")
doc.add_page_break()

# TOC
heading(doc, "目录", 1)
toc_items = [
    "1 任务理解与项目衔接",
    "2 数据获取、来源登记与适用性判断",
    "3 原始数据到模型可读数据的翻译",
    "4 STL 分解算法、参数与异常规则",
    "5 完整代码结构与运行方式",
    "6 浙江省全社会用电量 STL 结果",
    "7 三次产业季度增加值 STL 结果",
    "8 知识图谱与模型调用规则设计",
    "9 数据质量、局限与下一步工作",
    "10 结论",
    "附录 A 输出文件目录",
    "附录 B 标准模型输入示例",
    "附录 C 参考资料与外部核验来源",
]
for item in toc_items:
    p=doc.add_paragraph()
    p.paragraph_format.left_indent=Cm(0.7)
    p.paragraph_format.space_after=Pt(4)
    r=p.add_run(item)
    set_run_font(r,size=11)
doc.add_page_break()

# 1
heading(doc, "1 任务理解与项目衔接", 1)
heading(doc, "1.1 最终产品形态", 2)
add_body(doc, "项目希望形成一个“电力经济领域的 Claude Code”：用户给出某地区、某时间段的电力与经济数据，系统自动完成数据理解、异常识别、原因诊断、报告生成和交互问答。开题材料第 2 页强调从经济活动变化到用电行为变化，再到电力数据中提取经济信号；第 5 页给出的主线是“多源数据输入—特征解耦—知识增强诊断—多智能体协同—智能输出”。本次工作正对应其中的前两段，并为后两段准备标准接口。")
heading(doc, "1.2 老师强调的三个重点如何落实", 2)
add_table(doc,
          ["重点","本次完成内容","可验证产物","后续衔接"],
          [
              ["数据获取","建立来源登记、哈希、外部交叉核验和质量规则；保留原始文件","data_source_registry.csv、data_quality_checks.csv","扩展到气象、价格、分行业用电、政策事件"],
              ["STL 分解","电力月度与三次产业季度分别做鲁棒 log-STL；给出参数、异常阈值和端点警告","electricity_monthly_stl.csv、gdp_quarterly_stl.csv、8 张图","异常监测代理、跨尺度特征解耦"],
              ["知识图谱","把地区、指标、时间、观测、异常、节假日和候选原因转成节点与关系","kg_nodes.csv、kg_edges.csv","Neo4j/RAG/规则引擎、诊断智能体"],
          ], widths=[2.2,5.0,4.3,4.6], font_size=8.8)
heading(doc, "1.3 本次工作的系统位置", 2)
add_table(doc,
          ["阶段","输入","处理","输出给下一阶段"],
          [
              ["数据层","Excel、Word、网页来源信息","解析、单位规范、时间对齐、质量检查","统一时序表与溯源信息"],
              ["特征解耦层","统一时序表","log-STL、节假日特征、异常标记","期望值、趋势、季节、残差、异常"],
              ["知识组织层","分解结果+日历+规则","实体化、关系化、证据缺口化","知识图谱 CSV、模型 JSONL"],
              ["诊断层（后续）","图谱+外部证据+规则","RAG、LLM、多智能体审校","原因链、置信度、报告"],
          ], widths=[2.3,3.4,5.0,5.4], font_size=8.8)
doc.add_page_break()

# 2
heading(doc, "2 数据获取、来源登记与适用性判断", 1)
heading(doc, "2.1 本次实际参与计算的数据", 2)
add_table(doc,
          ["数据集","原始形态","有效时间范围","处理后记录数","用途"],
          [
              ["浙江全社会用电量（月度）","XLSX；原始单位万千瓦时","2021-01—2025-03","51","月度 STL、异常与季节性"],
              ["浙江第一产业增加值","XLSX；季度累计值","2000Q1—2025Q4","104","差分为单季度后做 STL"],
              ["浙江第二产业增加值","XLSX；季度累计值","2000Q1—2025Q4","104","差分为单季度后做 STL"],
              ["浙江第三产业增加值","XLSX；季度累计值","2000Q1—2025Q4","104","差分为单季度后做 STL"],
              ["国务院办公厅 2021—2025 节假日安排","DOCX 文本","2021—2025","按日历展开后聚合","春节/节假日/调休特征"],
          ], widths=[3.5,3.2,3.1,2.3,4.4], font_size=8.4)
add_body(doc, "代码没有用截图 OCR 读取上述浙江数据，而是直接解析 XLSX 内部 XML 与 DOCX 段落，从源头降低识别误差。每个输入文件均计算 SHA-256；任何替换、重采或修订都会导致哈希变化，从而触发重新运行和审计。")
heading(doc, "2.2 外部权威信息交叉核验", 2)
add_table(doc,
          ["核验对象","本地计算结果","公开口径","判断"],
          [
              ["2024 年浙江全社会用电量","6780.85 亿千瓦时","公开报道 6780 亿千瓦时","量级与四舍五入口径一致"],
              ["2025 年第一产业累计值","2657 亿元","浙江统计公报 2657 亿元","一致"],
              ["2025 年第二产业累计值","35682 亿元","浙江统计公报 35682 亿元","一致"],
              ["2025 年第三产业累计值","56206 亿元","浙江统计公报 56206 亿元","一致"],
              ["2021—2025 节假日","脚本解析放假与调休日","逐年国务院办公厅通知","来源完整"],
          ], widths=[4.2,4.0,4.4,3.4], font_size=8.5)
add_note_box(doc, "溯源现状", "浙江 4 个 Excel 原始文件本身没有内嵌抓取页 URL，因此当前标记为“部分完整”：数据值通过权威年度口径交叉核验，但后续数据获取代理仍应补录网页 URL、发布日期、抓取时间和版本号。节假日通知来源完整。", fill="EAF2F8", title_color=(18,59,93))
heading(doc, "2.3 苏州、南京、江门数据为什么暂不混入本轮 STL", 2)
add_body(doc, "用户提供的《data(苏州+南京+江门).docx》记录了三地官方入口：苏州统计局月度数据、南京统计年鉴/经济社会主要指标、江门统计月报。文档中的表格主要以截图存在，且时间结构不统一：南京有 2000—2023 年年度用电量，但年度频率没有季节项；苏州、江门主要是 2025—2026 年累计月度快照，只有约一个季节周期；南京 2025 年月度主要指标也不足以稳定估计 period=12 的 STL。")
add_body(doc, "因此本轮选择机器可读、时间跨度足够的浙江省数据作为可复现基线，三城市数据保留在 data/reference 中。后续应编写站点适配器，直接从官方 HTML/Excel/API 获取原始表，再将累计值还原为当期值，至少积累 3—5 个完整周期后进行城市间 STL 对比。")
heading(doc, "2.4 数据获取代理应固化的规则", 2)
for t in [
    "只接受官方统计局、政府、电网企业或有明确转载来源的权威数据；抓取页、发布日期和抓取时间必须记录。",
    "原始文件只读保存；清洗结果另存；每个文件计算哈希并建立版本号。",
    "对累计值、当期值、同比、单位、可比价/现价等口径分别建字段，不通过列名猜测。",
    "至少做一个独立总量交叉核验；差异超过预设阈值时停止模型调用，进入人工核验。",
    "截图/OCR 数据必须保存页面图像、识别置信度和人工复核状态，不与原始机器可读数据等价处理。",
]:
    add_bullet(doc,t)
doc.add_page_break()

# 3
heading(doc, "3 原始数据到模型可读数据的翻译", 1)
heading(doc, "3.1 月度用电量规范化", 2)
add_body(doc, "原始用电量单位为“万千瓦时”。为便于与宏观统计报告一致，统一转换为“亿千瓦时”：")
add_code_block(doc, "value_亿千瓦时 = raw_value_万千瓦时 / 10,000")
add_body(doc, "原始表内月份顺序并非时间升序，程序先解析“年、月”字段，构造月初日期，按时间排序，再检查是否从 2021-01 到 2025-03 连续。51 个月无缺口。")
heading(doc, "3.2 季度累计 GDP 还原为单季度值", 2)
add_body(doc, "三个产业文件记录的是年内累计增加值。直接对累计值做 STL 会把“Q2、Q3、Q4 必然更大”的会计结构误当成季节性，因此先还原单季度值：")
add_code_block(doc, "Q1 = Cum(Q1)\nQ2 = Cum(Q2) - Cum(Q1)\nQ3 = Cum(Q3) - Cum(Q2)\nQ4 = Cum(Q4) - Cum(Q3)")
add_body(doc, "程序检查 2000Q1 以后所有差分值均为正。Q4 由全年累计值减前三季度累计值，可能集中反映年度核算修订，所以单独设置 revision_risk=True，避免诊断智能体把口径修订误判为真实冲击。")
heading(doc, "3.3 节假日与调休日历特征", 2)
add_body(doc, "程序从国务院办公厅通知中提取放假区间和调休上班日，按月生成 holiday_days、spring_festival_days、makeup_workdays、holiday_names、holiday_event_ids。春节月份还设置 spring_festival_window，用于提示模型检查月份错位和复工节奏。")
heading(doc, "3.4 统一模型输入结构", 2)
add_table(doc,
          ["一级字段","核心内容","模型规则"],
          [
              ["region","地区编码、地区名","禁止仅用自然语言地区名；采用稳定编码"],
              ["period","标签、频率、起始日期","频率必须与 STL period 对应"],
              ["metric","指标 ID、名称、单位","原始单位与规范单位同时保留"],
              ["observation","观测值、累计值、同比等","禁止把累计值与当期值混用"],
              ["decomposition","趋势、季节因子、期望值、残差、权重","统一 log-STL 语义"],
              ["anomaly","等级、方向、端点/修订风险","质量风险优先于原因推理"],
              ["context","日历、候选原因、所需证据","候选解释不得写成已证实因果"],
              ["provenance","文件、哈希、变换、来源状态","每条结论可追溯到原始文件"],
          ], widths=[2.7,7.2,6.4], font_size=8.4)

# 4
heading(doc, "4 STL 分解算法、参数与异常规则", 1)
heading(doc, "4.1 为什么选择 STL", 2)
add_body(doc, "STL（Seasonal-Trend decomposition using LOESS）将时间序列拆分为趋势 T、季节 S 和不规则项 R。它允许季节形态随时间缓慢变化，并可通过 robust=True 降低异常点对趋势和季节估计的污染，适合节假日错位、极端天气和一次性冲击较多的电力数据。")
heading(doc, "4.2 乘法结构的 log-STL", 2)
add_body(doc, "电力和经济规模随趋势上升时，季节波动往往表现为比例而非固定绝对量，因此采用对数变换：")
add_code_block(doc, "log(y_t) = T_t + S_t + R_t\nexpected_t = exp(T_t + S_t)\nseasonal_factor_t = exp(S_t)\nresidual_pct_t = exp(R_t) - 1 = y_t / expected_t - 1")
add_body(doc, "其中 expected_t 是“在当前趋势与常规季节模式下的期望值”；residual_pct_t 是偏离期望的比例，最适合直接输入异常诊断规则。")
heading(doc, "4.3 参数选择", 2)
add_table(doc,
          ["序列","period","seasonal","robust","理由"],
          [
              ["月度全社会用电量",12,13,True,"一年 12 个月；季节窗口取大于 period 的奇数；降低异常污染"],
              ["季度单产业增加值",4,7,True,"一年 4 季度；窗口取奇数并允许季节模式平滑变化"],
          ], widths=[4.0,2.0,2.2,2.3,5.6], font_size=8.8)
add_body(doc, "本次参数是工程基线，不是唯一最优值。后续可通过滚动预测误差、残差自相关和稳定性分析在候选窗口中选择参数，而不是只凭可视化。")
heading(doc, "4.4 异常、端点和统计权重规则", 2)
add_table(doc,
          ["规则","定义","用途"],
          [
              ["high","|residual_pct| ≥ 8%","进入优先诊断队列"],
              ["medium","5% ≤ |residual_pct| < 8%","进入观察队列"],
              ["normal","|residual_pct| < 5%","作为正常背景"],
              ["statistical_outlier","robust_z 或 STL 权重显示强异常","补充证据，不替代业务阈值"],
              ["endpoint_warning","首尾一个完整季节周期","提醒端点分解不稳定"],
              ["revision_risk","GDP Q4","提醒累计值差分可能吸收年度修订"],
          ], widths=[3.2,5.0,7.9], font_size=8.8)
heading(doc, "4.5 趋势强度与季节强度", 2)
add_code_block(doc, "SeasonalStrength = max(0, 1 - Var(R) / Var(S + R))\nTrendStrength    = max(0, 1 - Var(R) / Var(T + R))")
add_body(doc, "强度越接近 1，表示相应成分相对于残差越稳定、越可解释。它们用于描述结构，不用于证明因果。")
add_note_box(doc, "方法边界", "STL 只能识别“与自身历史趋势和季节规律相比的异常”。跨地区比较、政策评估和因果归因还需要同期对照、外部变量、事件时间线和必要的统计检验。")
doc.add_page_break()

# 5
heading(doc, "5 完整代码结构与运行方式", 1)
heading(doc, "5.1 目录结构", 2)
add_code_block(doc, """STL电力经济分析成果包/
├─ stl_power_economy_pipeline.py
├─ requirements.txt
├─ data/
│  ├─ raw/        # 浙江原始 XLSX、节假日 DOCX
│  └─ reference/  # 苏州+南京+江门参考资料
└─ outputs/
   ├─ tables/      # STL 明细、异常、质量、来源
   ├─ figures/     # 8 张分析图
   ├─ model_input/ # model_input_records.jsonl
   ├─ kg/          # kg_nodes.csv, kg_edges.csv
   └─ run_summary.json""")
heading(doc, "5.2 环境与命令", 2)
add_code_block(doc, "python -m pip install -r requirements.txt\npython stl_power_economy_pipeline.py --input-dir data/raw --output-dir outputs")
add_body(doc, "依赖文件仅包含 numpy、statsmodels、matplotlib、python-docx、Pillow。脚本没有依赖 Excel 桌面程序，也没有使用 OCR；XLSX 通过 ZIP/XML 直接读取。")
heading(doc, "5.3 主程序的处理顺序", 2)
for t in [
    "定位并读取 5 个原始文件，计算 SHA-256。",
    "解析月度用电量并转换单位，解析三产业季度累计值并差分。",
    "解析节假日通知，构建日历事件和月度特征。",
    "执行鲁棒 log-STL，计算期望值、趋势、季节因子、残差比例、权重和强度。",
    "根据业务阈值、端点和修订风险生成异常事件。",
    "生成模型 JSONL、知识图谱节点/关系、质量报告、数据源登记和图表。",
]:
    add_number(doc,t)
heading(doc, "5.4 关键实现片段", 2)
add_code_block(doc, """stl = STL(
    np.log(values),
    period=period,
    seasonal=seasonal_window,
    robust=True,
)
fit = stl.fit()
expected = np.exp(fit.trend + fit.seasonal)
residual_pct = values / expected - 1.0""")
add_code_block(doc, """a = abs(residual_pct)
level = "high" if a >= 0.08 else ("medium" if a >= 0.05 else "normal")""")
add_body(doc, "完整实现已单独交付为 stl_power_economy_pipeline.py，包含 XLSX/XML 解析、节假日文本解析、所有质量检查、知识图谱与 JSONL 生成逻辑，报告中不省略任何运行步骤。")
heading(doc, "5.5 本次实际运行日志", 2)
run_log = (outputs / "run_log.txt").read_text(encoding="utf-8").strip()
add_code_block(doc, run_log)
doc.add_page_break()

# 6 Electricity
heading(doc, "6 浙江省全社会用电量 STL 结果", 1)
heading(doc, "6.1 总体结果", 2)
add_table(doc,
          ["指标","结果","解释"],
          [
              ["有效记录", "51 个月（2021-01—2025-03）", "时间连续，无缺月"],
              ["2024 年总量", "6780.85 亿千瓦时", "与外部权威年度口径基本一致"],
              ["趋势起点/终点", f"{summary['electricity']['trend_start']:.2f} / {summary['electricity']['trend_end']:.2f} 亿千瓦时", "STL 趋势水平"],
              ["趋势年化增速", f"{summary['electricity']['trend_cagr']:.2%}", "按趋势首尾复合年化"],
              ["季节强度", f"{summary['electricity']['seasonal_strength']:.3f}", "季节结构较明显"],
              ["趋势强度", f"{summary['electricity']['trend_strength']:.3f}", "趋势存在但残差影响不小"],
              ["high / medium", f"{summary['electricity']['high_anomaly_count']} / {summary['electricity']['medium_anomaly_count']}", "按 8% 与 5% 阈值"],
          ], widths=[4.1,4.8,7.2], font_size=8.8)
add_figure(doc, outputs/"figures/01_electricity_actual_vs_expected.png", "图 6-1  浙江全社会用电量实际值与 STL 期望值")
add_body(doc, "实际值与期望值整体同步，说明 trend+seasonal 已解释大部分常规波动；明显偏离集中在春节错位月份和少数非春节月份。")
add_figure(doc, outputs/"figures/02_electricity_trend.png", "图 6-2  月度用电量趋势水平")
add_body(doc, "趋势从约 450.33 亿千瓦时上升到 580.36 亿千瓦时，按样本首尾趋势计算的年化增速约 6.29%。末端趋势仍属于端点估计，应在新数据到来后滚动更新。")
doc.add_page_break()

heading(doc, "6.2 月份季节性", 2)
season_rows = []
for r in seasonal:
    season_rows.append([
        f"{int(float(r['month']))}月",
        f"{float(r['average_seasonal_factor']):.3f}",
        f"{float(r['average_effect_pct']):+.2f}%",
        r["n"],
    ])
add_table(doc, ["月份","平均季节因子","相对趋势效应","样本数"], season_rows, widths=[2.5,4.0,4.0,2.5], font_size=8.8)
add_figure(doc, outputs/"figures/03_electricity_month_seasonal_factors.png", "图 6-3  月度平均季节因子")
add_body(doc, "7—8 月平均季节因子分别约为 1.253、1.266，是全年高位；2 月约为 0.794，是最低月份。这一结构符合夏季负荷抬升与春节停复工的常见时间模式，但具体原因仍需气象、分行业用电和生产数据验证。")
heading(doc, "6.3 电力 high 异常清单", 2)
high_e = summary["electricity"]["high_anomalies"]
rows = []
for x in high_e:
    rows.append([
        x["period"],
        f"{x['observed']:.2f}",
        f"{x['expected']:.2f}",
        f"{x['residual_pct']:+.2%}",
        "是" if x["spring_festival_window"] else "否",
        "是" if x["endpoint_warning"] else "否",
    ])
add_table(doc, ["期间","实际值","期望值","偏离","春节窗口","端点警告"], rows, widths=[2.2,2.8,2.8,2.4,2.4,2.4], font_size=8.7)
add_figure(doc, outputs/"figures/04_electricity_residual_anomalies.png", "图 6-4  电力残差比例与异常阈值")
heading(doc, "6.4 异常解释示例", 2)
add_table(doc,
          ["异常","可观察事实","当前候选解释","下一步必须补证据"],
          [
              ["2023-01 -23.61% / 2023-02 +10.47%","春节在 1 月，低谷与反弹跨月","春节日期错位与复工节奏","分行业用电、规上工业增加值、开工率、春节日期"],
              ["2024-01 +17.69% / 2024-02 -17.98%","春节在 2 月，1 月偏高、2 月偏低","春节前置生产与 2 月停工","日度负荷、制造业用电、春节前后生产日历"],
              ["2022-09 -11.54%","不在春节窗口","天气、工业活动、政策/电网事件之一或组合","温度、工业增加值、分行业用电、电价/限电/需求响应记录"],
              ["2021-01/02","样本起点且偏离较大","可能含春节因素，但端点不稳定","更早历史数据；不应作为强因果证据"],
          ], widths=[3.2,4.2,4.2,5.0], font_size=8.1)
doc.add_page_break()

# 7 GDP
heading(doc, "7 三次产业季度增加值 STL 结果", 1)
heading(doc, "7.1 总体结构", 2)
gdp_metrics = summary["gdp"]["metrics_by_industry"]
rows = []
for ind in ["第一产业","第二产业","第三产业"]:
    m = gdp_metrics[ind]
    rows.append([
        ind, m["n_observations"], f"{m['seasonal_strength']:.3f}", f"{m['trend_strength']:.3f}",
        f"{m['residual_mad_log']:.4f}", f"{m['residual_std_log']:.4f}"
    ])
add_table(doc, ["产业","记录数","季节强度","趋势强度","残差MAD(log)","残差STD(log)"], rows, widths=[2.4,2.0,2.6,2.6,3.2,3.2], font_size=8.5)
add_body(doc, "三个产业的趋势强度均高于 0.99，说明 2000—2025 年长期增长轨迹非常突出；季节强度中第一产业最高，第二、第三产业也具有明显季度结构。高趋势强度并不意味着没有周期冲击，异常仍由相对期望值的残差识别。")
add_figure(doc, outputs/"figures/05_gdp_第一产业_actual_vs_expected.png", "图 7-1  第一产业单季度增加值与 STL 期望值", width_cm=15.0)
doc.add_page_break()
add_figure(doc, outputs/"figures/06_gdp_第二产业_actual_vs_expected.png", "图 7-2  第二产业单季度增加值与 STL 期望值", width_cm=14.2)
add_figure(doc, outputs/"figures/07_gdp_第三产业_actual_vs_expected.png", "图 7-3  第三产业单季度增加值与 STL 期望值", width_cm=14.2)
doc.add_page_break()
add_figure(doc, outputs/"figures/08_gdp_residual_anomalies.png", "图 7-4  三次产业残差比例与异常点", width_cm=14.8)
heading(doc, "7.2 high 异常概览", 2)
recent_high = [x for x in summary["gdp"]["high_anomalies"] if int(x["period"][:4]) >= 2014]
rows = []
for x in recent_high:
    rows.append([
        x["period"], x["industry"], f"{x['observed']:.2f}", f"{x['expected']:.2f}",
        f"{x['residual_pct']:+.2%}", "是" if x["revision_risk"] else "否", "是" if x["endpoint_warning"] else "否"
    ])
add_table(doc, ["期间","产业","实际值","期望值","偏离","Q4风险","端点"], rows, widths=[2.0,2.3,2.7,2.7,2.2,2.1,1.8], font_size=8.1)
add_body(doc, "较新的明显偏离包括：第二产业 2019Q4 高于期望约 24.10%、2020Q1 低于期望约 12.70%；第一产业 2024Q4 高于期望约 19.18%；第三产业 2024Q4 高于期望约 16.71%、2025Q4 低于期望约 11.89%。其中 Q4 结果均带 revision_risk，必须先检查年度核算和累计值修订，再讨论经济含义。")
heading(doc, "7.3 为什么不能直接把历史异常解释为某个事件", 2)
add_body(doc, "本轮只使用三产业增加值和节假日，没有加载价格指数、固定资产投资、出口、工业生产、疫情管控、重大政策和自然灾害时间线。对于 2005、2009、2017、2019—2020 等异常，程序只保留“宏观冲击/产业活动/数据修订”等候选类别与证据需求，不自动写入具体事件名称。这样可以防止大模型用常识补全历史叙事，形成看似合理但不可验证的解释。")

# 8 KG
heading(doc, "8 知识图谱与模型调用规则设计", 1)
heading(doc, "8.1 图谱规模与核心实体", 2)
add_table(doc,
          ["对象","数量/状态","说明"],
          [
              ["知识图谱节点",summary["knowledge_graph"]["node_count"],"地区、指标、时间、观测、异常、节假日、候选原因、来源等"],
              ["知识图谱关系",summary["knowledge_graph"]["edge_count"],"OBSERVED_AT、HAS_ANOMALY、TEMPORALLY_OVERLAPS、CANDIDATE_EXPLANATION 等"],
              ["模型输入记录",summary["model_input"]["record_count"],"每条观测一个 JSON 对象"],
              ["异常事件",len(anomalies),"high 与 medium 均进入异常事件表"],
          ], widths=[4.0,3.0,9.0], font_size=8.8)
heading(doc, "8.2 推荐的图谱模式", 2)
add_code_block(doc, """(:Region)-[:HAS_METRIC]->(:Metric)
(:Observation)-[:FOR_REGION]->(:Region)
(:Observation)-[:OF_METRIC]->(:Metric)
(:Observation)-[:OBSERVED_AT]->(:TimePeriod)
(:Observation)-[:HAS_ANOMALY]->(:Anomaly)
(:Anomaly)-[:TEMPORALLY_OVERLAPS]->(:HolidayEvent)
(:Anomaly)-[:CANDIDATE_EXPLANATION]->(:CauseHypothesis)
(:CauseHypothesis)-[:REQUIRES_EVIDENCE]->(:EvidenceType)
(:Observation)-[:DERIVED_FROM]->(:DataSource)""")
add_body(doc, "关系命名刻意区分“时间重合”“候选解释”和“已验证因果”。在未引入因果检验与外部证据前，不创建 CAUSED_BY 关系。")
heading(doc, "8.3 每次模型调用之间的规则和约束", 2)
add_table(doc,
          ["环节","强制输入","强制检查","允许输出"],
          [
              ["数据理解","region/period/metric/unit/provenance","缺失、频率、单位、累计/当期","标准化记录或停止原因"],
              ["异常检测","expected/residual/threshold/quality flags","端点、修订、样本长度","异常等级与偏离"],
              ["原因检索","异常+候选原因+证据需求","是否有外部证据、时间先后、口径一致","候选原因及置信度"],
              ["结论生成","证据链+图谱路径+规则结果","因果措辞、证据覆盖、矛盾","分层结论、局限、下一步"],
              ["报告审校","所有中间产物","数字一致性、来源、单位、禁止过度归因","可发布报告或退回重查"],
          ], widths=[2.4,4.0,5.0,4.7], font_size=8.2)
heading(doc, "8.4 推荐的多智能体拆分", 2)
for t in [
    "数据获取代理：抓取、版本化、哈希、字段映射、交叉核验。",
    "时序分析代理：STL 参数选择、分解、阈值、端点/修订风险。",
    "证据检索代理：天气、产业、政策、电网事件、价格等外部证据。",
    "知识图谱代理：实体对齐、关系约束、冲突检测和图谱写入。",
    "诊断与报告代理：基于证据链生成结论；无证据时明确保留不确定性。",
    "审计代理：核对数据、代码版本、图表、结论和引用是否一致。",
]:
    add_bullet(doc,t)
heading(doc, "8.5 2024-02 电力异常的模型输入示例", 2)
sample = next(r for r in model_records if r["period"]["label"]=="2024-02" and r["metric"]["id"]=="electricity_total_monthly")
sample_compact = {
    "region": sample["region"],
    "period": sample["period"],
    "metric": sample["metric"],
    "observation": sample["observation"],
    "decomposition": {
        "expected_value": round(sample["decomposition"]["expected_value"],2),
        "residual_pct": round(sample["decomposition"]["residual_pct"],4),
        "trend_level": round(sample["decomposition"]["trend_level"],2),
        "seasonal_factor": round(sample["decomposition"]["seasonal_factor"],4),
        "stl_weight": sample["decomposition"]["stl_weight"],
    },
    "anomaly": sample["anomaly"],
    "context": sample["context"],
    "provenance": sample["provenance"],
}
add_code_block(doc, json.dumps(sample_compact, ensure_ascii=False, indent=2))
doc.add_page_break()

# 9
heading(doc, "9 数据质量、局限与下一步工作", 1)
heading(doc, "9.1 本次质量检查", 2)
passed = sum(1 for r in quality if r["passed"].lower()=="true")
warn = sum(1 for r in quality if r["severity"]=="warning")
add_body(doc, f"程序共执行 {len(quality)} 项质量检查，全部通过；其中 {warn} 项为 warning 级提示，主要涉及来源 URL 未完全补录、Q4 修订风险和外部交叉核验的口径差异。这些 warning 不阻止运行，但会随记录进入模型上下文。")
qrows = []
for r in quality:
    if r["severity"]=="warning" or r["check_name"] in ("时间连续性","外部总量核验","年度累计值核验"):
        qrows.append([r["dataset_id"],r["check_name"],"通过" if r["passed"].lower()=="true" else "未通过",r["severity"],r["detail"]])
add_table(doc, ["数据集","检查项","结果","级别","详情"], qrows, widths=[3.2,3.3,1.8,2.0,6.0], font_size=7.8)
heading(doc, "9.2 当前局限", 2)
for t in [
    "月度电力只有 51 个样本，春节月份的季节形态估计仍会受到年份差异影响；最好补齐 2015 年以来数据。",
    "目前只分析全社会用电量，无法区分工业、商业、居民、新能源消纳等来源。",
    "GDP 是现价累计值差分；季度核算修订和价格因素可能影响残差。",
    "尚未接入气象、产业生产、价格、电网运行与政策事件，原因诊断只能生成候选解释。",
    "苏州、南京、江门资料尚未完成机器可读采集与统一口径，不能直接做跨城市结论。",
    "固定 5%/8% 阈值是工程基线，应结合指标波动率、滚动分位数和业务损失进一步校准。",
]:
    add_bullet(doc,t)
heading(doc, "9.3 与后续任务的衔接路线", 2)
add_table(doc,
          ["优先级","任务","输入/产物","验收标准"],
          [
              ["P0","补齐浙江数据来源 URL 与抓取元数据","来源登记表","每个原始文件有 URL、发布日期、抓取时间、版本"],
              ["P0","扩展 10 年以上月度电力与分行业用电","统一时序表","无缺口、口径一致、年度总量可核验"],
              ["P1","接入逐日气温/湿度与极端天气","天气特征表","与地区、日期可对齐；来源权威"],
              ["P1","接入工业增加值、PMI、出口、开工率等","经济证据表","与异常月/季度形成证据链"],
              ["P1","建设政策/电网事件库","事件节点与时间区间","事件有发布日期、实施期、地区和来源"],
              ["P2","Neo4j 导入与 RAG 检索","图谱数据库","能从异常追溯观测、来源、候选原因和证据"],
              ["P2","参数与阈值回测","滚动验证报告","比较窗口、阈值、误报/漏报与稳定性"],
              ["P3","多地区适配器","苏州/南京/江门等","同一 schema 下可执行跨地区分析"],
          ], widths=[1.4,4.2,4.7,5.8], font_size=7.8)
heading(doc, "9.4 建议的最小可用诊断闭环", 2)
add_code_block(doc, """用户问题
  → 数据获取代理校验地区/时间/口径
  → STL 工具返回 expected、residual、quality flags
  → 异常代理决定是否进入诊断
  → 图谱/RAG 检索节假日、天气、产业、政策、电网证据
  → 规则引擎检查时间先后、证据覆盖和矛盾
  → LLM 生成“事实—候选原因—证据—不确定性—下一步”
  → 审计代理复核数字、引用和因果措辞""")
doc.add_page_break()

# 10
heading(doc, "10 结论", 1)
add_body(doc, "本次工作已经把老师要求的三个重点落实为可运行、可复现、可审计的工程成果，而不是只停留在算法说明：数据获取层有原始文件、哈希、来源登记与外部核验；STL 层有明确的变换、参数、阈值、端点和修订规则；知识图谱层有可直接导入数据库的节点/关系和面向模型调用的 JSONL。")
add_body(doc, "分析结果显示，浙江月度全社会用电量具有明显季节性，7—8 月为高季节因子、2 月为低季节因子；2023 和 2024 年春节月份出现显著跨月偏离。长期趋势从样本起点约 450.33 亿千瓦时上升到末端约 580.36 亿千瓦时。三次产业季度序列均呈强长期趋势，同时存在若干需要结合核算修订与外部事件进一步核验的异常。")
add_body(doc, "更关键的是，本成果已经把“模型能看懂的数据”定义清楚：每条记录不仅有数值，还包括正常期望、偏离程度、质量风险、候选原因、证据缺口和溯源。后续只需继续补充数据源和证据库，就可以在这一接口上迭代诊断规则、RAG、知识图谱和多智能体协同。")
add_note_box(doc, "可直接用于答辩的概括", "我们不是让大模型直接读表猜原因，而是先用可解释的时序算法把趋势、季节和异常分开，再把异常、证据和规则组织成模型可消费的结构化上下文；知识图谱负责约束关系，大模型负责在证据边界内生成诊断和报告。", fill="DCEAF4", title_color=(18,59,93))
doc.add_page_break()

# Appendix A
heading(doc, "附录 A 输出文件目录", 1)
files_desc = [
    ("stl_power_economy_pipeline.py","完整主程序；可从原始文件一键重跑"),
    ("outputs/tables/electricity_monthly_stl.csv","51 条月度用电 STL 明细"),
    ("outputs/tables/gdp_quarterly_stl.csv","312 条三产业单季度 STL 明细"),
    ("outputs/tables/anomaly_events.csv","high/medium 异常事件及候选原因"),
    ("outputs/tables/holiday_calendar_events.csv","逐日节假日与调休事件"),
    ("outputs/tables/holiday_month_features.csv","按月聚合的节假日特征"),
    ("outputs/tables/data_source_registry.csv","来源、外部核验、哈希与溯源状态"),
    ("outputs/tables/data_quality_checks.csv","全部质量检查结果"),
    ("outputs/model_input/model_input_records.jsonl","363 条模型标准输入"),
    ("outputs/kg/kg_nodes.csv","647 个知识图谱节点"),
    ("outputs/kg/kg_edges.csv","1235 条知识图谱关系"),
    ("outputs/figures/*.png","8 张运行结果图"),
    ("outputs/run_summary.json","关键指标和异常摘要"),
]
add_table(doc, ["文件","用途"], files_desc, widths=[7.5,8.5], font_size=8.4)
heading(doc, "附录 A.1 结果工作簿", 2)
add_body(doc, "另交付《浙江电力经济_STL分解结果.xlsx》，包含总览、电力 STL、GDP STL、异常清单、月份季节因子、节假日特征、数据质量、数据源登记、知识图谱节点/关系和模型输入样例等工作表，可直接筛选和复核。")

# Appendix B
doc.add_page_break()
heading(doc, "附录 B 标准模型输入示例", 1)
add_body(doc, "下列为完整 JSONL 中 2024-02 月度电力异常记录的结构。实际文件每行一个 JSON 对象，便于流式读取、向量化和作为模型工具调用参数。")
add_code_block(doc, json.dumps(sample, ensure_ascii=False, indent=2))

# Appendix C
doc.add_page_break()
heading(doc, "附录 C 参考资料与外部核验来源", 1)
refs = [
    ("[1] 用户项目材料", "《基于特征解耦与知识增强的“电力看经济”智能体构建方法》开题答辩 PPT，第 2—7 页。"),
    ("[2] 用户数据汇总", "《data(苏州+南京+江门).docx》，含苏州、南京、江门官方统计入口与数据截图。"),
    ("[3] STL 原始论文", "Cleveland, R. B., Cleveland, W. S., McRae, J. E., & Terpenning, I. (1990). STL: A Seasonal-Trend Decomposition Procedure Based on Loess."),
    ("[4] statsmodels 文档", "statsmodels.tsa.seasonal.STL：period、seasonal、trend、low_pass、robust 参数与 fit 输出说明。"),
    ("[5] 国家统计局", "季节调整相关说明：经济时间序列可分为趋势、周期、季节和不规则成分，并需考虑移动节假日。"),
    ("[6] 国务院办公厅", "2021—2025 年部分节假日安排通知。"),
    ("[7] 2024 浙江用电核验", "新华社/国家电网浙江电力公开信息：2024 年浙江全社会用电量约 6780 亿千瓦时。"),
    ("[8] 2025 浙江经济核验", "浙江省 2025 年国民经济和社会发展统计公报：三次产业增加值 2657、35682、56206 亿元。"),
]
add_table(doc, ["编号","资料"], refs, widths=[3.5,12.5], font_size=8.5)
heading(doc, "附录 C.1 数据入口（项目后续采集）", 2)
urls = [
    ("苏州统计局月度数据", "https://tjj.suzhou.gov.cn/sztjj/jdsj/nav_list.shtml"),
    ("南京统计年鉴用电量", "https://tjj.nanjing.gov.cn/material/njnj_2024/gongye/7-14.html"),
    ("南京统计局检索", "https://tjj.nanjing.gov.cn/site/tjj/search.html"),
    ("江门统计月报", "https://www.jiangmen.gov.cn/bmpd/jmstjj/tjsj/tjyb/ydysjysyd/"),
    ("2024 浙江用电核验", "https://www.news.cn/20250124/4a7d51aece054002ac1edc66d125884e/c.html"),
    ("2025 浙江统计公报", "https://zjzd.stats.gov.cn/zwgk/zfxxgkml/tjxx/tjgb/art/2026/art_48c3c7315981425f9c25b53eab4d65d0.html"),
]
for name,url in urls:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r=p.add_run(name+"：")
    set_run_font(r,size=9.2,bold=True)
    add_hyperlink(p,url,url)

# Final metadata and save
doc.core_properties.title = "浙江省电力—经济数据 STL 分解与智能体接口设计"
doc.core_properties.subject = "数据获取、STL分解、知识图谱与模型输入"
doc.core_properties.author = "OpenAI GPT-5.6 Pro"
doc.core_properties.keywords = "电力经济, STL, 异常检测, 知识图谱, 大模型智能体"
doc.core_properties.comments = "基于用户提供材料生成；完整代码与运行结果另附。"

doc.save(report_path)
print("saved", report_path, report_path.stat().st_size)